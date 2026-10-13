"""Lectura y escritura de archivos Excel (.xlsx, .xls, .xlsm) y CSV con autodetección y formato."""

from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

from backend.core.config import ensure_output_dir
from backend.core.utils import FriendlyAppError
from backend.core.worker import CancelToken

SUPPORTED_EXCEL_EXTS: set[str] = {".xlsx", ".xls", ".xlsm"}
SUPPORTED_CSV_EXTS: set[str] = {".csv", ".tsv", ".txt"}
SUPPORTED_ALL_EXTS: set[str] = SUPPORTED_EXCEL_EXTS | SUPPORTED_CSV_EXTS

CANDIDATE_DELIMITERS: list[str] = [",", ";", "\t", "|"]
CANDIDATE_ENCODINGS: list[str] = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]


def clean_file_path(raw_path: str | Path) -> Path:
    """Limpia rutas provenientes de QML (file:///) y devuelve un objeto Path."""
    from urllib.parse import unquote

    s = str(raw_path).strip()
    if s.startswith("file:///"):
        s = s[8:]
    elif s.startswith("file://"):
        s = s[7:]
    return Path(unquote(s))


def detect_csv_encoding(file_path: Path | str, sample_bytes: int = 65536) -> str:
    """
    Detecta automáticamente la codificación de un archivo CSV/TSV/TXT.
    Conserva correctamente acentos y la letra ñ.
    """
    path = clean_file_path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"No existe el archivo: {path}")

    with open(path, "rb") as f:
        raw = f.read(sample_bytes)

    if not raw:
        return "utf-8-sig"

    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"

    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        pass

    try:
        raw.decode("cp1252")
        return "cp1252"
    except UnicodeDecodeError:
        return "latin-1"


def detect_csv_delimiter(
    file_path: Path | str,
    encoding: str | None = None,
    sample_lines: int = 30,
) -> str:
    """
    Detecta automáticamente el delimitador de un archivo CSV (',', ';', '\\t', '|').
    """
    path = clean_file_path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"No existe el archivo: {path}")

    enc = encoding if (encoding and encoding != "auto") else detect_csv_encoding(path)

    lines: list[str] = []
    with open(path, "r", encoding=enc, errors="replace") as f:
        for _ in range(sample_lines):
            line = f.readline()
            if not line:
                break
            stripped = line.strip("\r\n")
            if stripped.strip():
                lines.append(stripped)

    if not lines:
        return ";" if path.suffix.lower() == ".csv" else ","

    if path.suffix.lower() == ".tsv" and any("\t" in ln for ln in lines):
        return "\t"

    # Evaluar consistencia de columnas por cada delimitador candidato
    best_delim = ","
    best_score = -1.0

    for delim in CANDIDATE_DELIMITERS:
        counts: list[int] = []
        for ln in lines[:15]:
            try:
                row = next(csv.reader([ln], delimiter=delim))
                counts.append(len(row))
            except Exception:
                counts.append(1)

        if not counts or max(counts) <= 1:
            continue

        first_count = counts[0]
        if first_count <= 1:
            continue

        consistency = sum(1 for c in counts if c == first_count) / len(counts)
        score = consistency * 100.0 + min(first_count, 50)
        if score > best_score:
            best_score = score
            best_delim = delim

    if best_score > 0:
        return best_delim

    # Fallback con csv.Sniffer
    try:
        sample_text = "\n".join(lines[:10])
        dialect = csv.Sniffer().sniff(sample_text, delimiters=",;\t|")
        return dialect.delimiter
    except Exception:
        return ","


def inspect_file_metadata(file_path: Path | str) -> dict[str, Any]:
    """
    Inspecciona rápidamente un archivo para devolver su tipo, hojas (si es Excel)
    o delimitador/codificación detectados (si es CSV).
    """
    path = clean_file_path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"No se encontró el archivo: {path}")

    ext = path.suffix.lower()
    size_bytes = path.stat().st_size
    if size_bytes == 0:
        raise FriendlyAppError("El archivo seleccionado está vacío (0 bytes).")

    if ext in SUPPORTED_EXCEL_EXTS:
        sheets = get_excel_sheets(path)
        return {
            "path": str(path),
            "name": path.name,
            "isExcel": True,
            "isCsv": False,
            "sheets": sheets,
            "defaultSheet": sheets[0] if sheets else "",
            "detectedEncoding": "",
            "detectedDelimiter": "",
            "sizeBytes": size_bytes,
        }

    enc = detect_csv_encoding(path)
    delim = detect_csv_delimiter(path, encoding=enc)
    return {
        "path": str(path),
        "name": path.name,
        "isExcel": False,
        "isCsv": True,
        "sheets": [],
        "defaultSheet": "",
        "detectedEncoding": enc,
        "detectedDelimiter": delim,
        "sizeBytes": size_bytes,
    }


def get_excel_sheets(file_path: Path | str) -> list[str]:
    """Obtiene los nombres de las hojas de un archivo Excel."""
    path = clean_file_path(file_path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXCEL_EXTS:
        return []

    if ext in {".xlsx", ".xlsm"}:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()

    # Para .xls intentamos pd.ExcelFile; si es un TSV/CSV con extensión .xls devolvemos ["Hoja1"]
    try:
        with pd.ExcelFile(path) as xls:
            return [str(s) for s in xls.sheet_names]
    except Exception:
        return ["Hoja1"]


def normalize_delimiter_input(delim: str | None) -> str | None:
    """Convierte etiquetas de UI o secuencias escapadas en el carácter delimitador real."""
    if not delim or delim.lower() in {"auto", "automático", "automatico"}:
        return None
    mapping = {
        ",": ",",
        "coma (,)": ",",
        ";": ";",
        "punto y coma (;)": ";",
        "\\t": "\t",
        "\t": "\t",
        "tab": "\t",
        "tabulador (\\t)": "\t",
        "|": "|",
        "barra vertical (|)": "|",
    }
    return mapping.get(delim.strip().lower(), delim if len(delim) == 1 else ",")


def normalize_encoding_input(enc: str | None) -> str | None:
    """Convierte etiquetas de UI en nombres de codificación estándar de Python."""
    if not enc or enc.lower() in {"auto", "automático", "automatico"}:
        return None
    mapping = {
        "utf-8": "utf-8",
        "utf-8-sig": "utf-8-sig",
        "utf-8 con bom (utf-8-sig)": "utf-8-sig",
        "latin-1": "latin-1",
        "cp1252": "cp1252",
        "latin-1 / cp1252": "cp1252",
    }
    return mapping.get(enc.strip().lower(), enc.strip())


def read_tabular_file(
    file_path: Path | str,
    sheet_name: str | int | None = None,
    delimiter: str | None = None,
    encoding: str | None = None,
    infer_types: bool = False,
    nrows: int | None = None,
    chunksize: int = 25000,
    progress_callback: Callable[[int, str], None] | None = None,
    cancel_token: CancelToken | None = None,
) -> pd.DataFrame:
    """
    Lee un archivo Excel (.xlsx, .xls, .xlsm) o delimitado (.csv, .tsv, .txt) en un DataFrame.

    Por defecto (`infer_types=False`) lee las columnas como texto (`dtype=str`) para no perder
    ceros a la izquierda en códigos o SKUs, manteniendo celdas vacías como NaN/None.
    """
    path = clean_file_path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"No se encontró el archivo: {path}")

    if path.stat().st_size == 0:
        raise FriendlyAppError("El archivo seleccionado está vacío (0 bytes).")

    if cancel_token:
        cancel_token.raise_if_cancelled()

    ext = path.suffix.lower()
    dtype_arg: Any = None if infer_types else str

    if ext in SUPPORTED_EXCEL_EXTS:
        if progress_callback:
            progress_callback(15, f"Leyendo libro Excel: {path.name}...")

        target_sheet: str | int = 0
        if sheet_name is not None and str(sheet_name).strip() != "":
            target_sheet = sheet_name

        try:
            if ext in {".xlsx", ".xlsm"}:
                df = pd.read_excel(
                    path,
                    sheet_name=target_sheet,
                    dtype=dtype_arg,
                    nrows=nrows,
                    engine="openpyxl",
                )
            else:
                try:
                    df = pd.read_excel(
                        path,
                        sheet_name=target_sheet,
                        dtype=dtype_arg,
                        nrows=nrows,
                    )
                except Exception:
                    # Fallback para archivos .xls exportados como texto delimitado por ERP
                    enc = detect_csv_encoding(path)
                    delim = detect_csv_delimiter(path, encoding=enc)
                    df = pd.read_csv(
                        path,
                        sep=delim,
                        encoding=enc,
                        dtype=dtype_arg,
                        nrows=nrows,
                        low_memory=False,
                    )
        except ValueError as exc:
            if "Worksheet" in str(exc) or "not found" in str(exc):
                raise FriendlyAppError(
                    f"La hoja '{target_sheet}' no existe en el archivo {path.name}.",
                    str(exc),
                ) from exc
            raise

        if cancel_token:
            cancel_token.raise_if_cancelled()

        if progress_callback:
            progress_callback(85, f"Archivo Excel leído ({len(df):,} filas).")

        return df

    # Lectura de CSV / TSV / TXT
    eff_enc = normalize_encoding_input(encoding) or detect_csv_encoding(path)
    eff_delim = normalize_delimiter_input(delimiter) or detect_csv_delimiter(path, encoding=eff_enc)

    if nrows is not None:
        try:
            return pd.read_csv(
                path,
                sep=eff_delim,
                encoding=eff_enc,
                dtype=dtype_arg,
                nrows=nrows,
                low_memory=False,
            )
        except UnicodeDecodeError:
            return pd.read_csv(
                path,
                sep=eff_delim,
                encoding="latin-1",
                dtype=dtype_arg,
                nrows=nrows,
                low_memory=False,
            )

    # Lectura por bloques (chunks) para progreso real y cancelación cooperativa
    if progress_callback:
        progress_callback(10, f"Iniciando lectura de {path.name} ({eff_enc}, sep='{eff_delim}')...")

    # Estimar líneas totales rápidamente por bloques binarios
    total_lines = 0
    with open(path, "rb") as bf:
        while chunk_bytes := bf.read(1024 * 1024):
            total_lines += chunk_bytes.count(b"\n")
    total_data_rows = max(1, total_lines - 1)

    chunks: list[pd.DataFrame] = []
    rows_read = 0

    def _iter_csv(enc_to_use: str):
        return pd.read_csv(
            path,
            sep=eff_delim,
            encoding=enc_to_use,
            dtype=dtype_arg,
            chunksize=chunksize,
            low_memory=False,
        )

    try:
        reader = _iter_csv(eff_enc)
        for chunk in reader:
            if cancel_token:
                cancel_token.raise_if_cancelled()
            chunks.append(chunk)
            rows_read += len(chunk)
            if progress_callback:
                pct = min(85, 10 + int((rows_read / total_data_rows) * 75))
                progress_callback(pct, f"Leyendo CSV: {rows_read:,} filas procesadas...")
    except UnicodeDecodeError:
        chunks.clear()
        rows_read = 0
        reader = _iter_csv("latin-1")
        for chunk in reader:
            if cancel_token:
                cancel_token.raise_if_cancelled()
            chunks.append(chunk)
            rows_read += len(chunk)
            if progress_callback:
                pct = min(85, 10 + int((rows_read / total_data_rows) * 75))
                progress_callback(pct, f"Leyendo CSV (latin-1): {rows_read:,} filas procesadas...")

    if not chunks:
        # Archivo solo con encabezados y 0 filas de datos
        return pd.read_csv(
            path,
            sep=eff_delim,
            encoding=eff_enc,
            dtype=dtype_arg,
            nrows=0,
        )

    df = pd.concat(chunks, ignore_index=True)
    return df


def build_output_path(
    source_path_or_name: Path | str,
    suffix_tag: str,
    extension: str,
    output_dir: Path | str | None = None,
) -> Path:
    """
    Construye una ruta de archivo nueva con sufijo descriptivo y fecha/hora,
    garantizando que NUNCA se sobrescriba un archivo existente.
    """
    out_folder = ensure_output_dir(output_dir)
    src_stem = clean_file_path(source_path_or_name).stem or "resultado"
    # Limpiar caracteres inválidos en nombres de archivo de Windows
    safe_stem = "".join(c if c not in '<>:"/\\|?*' else "_" for c in src_stem).strip()
    safe_tag = "".join(c if c not in '<>:"/\\|?*' else "_" for c in suffix_tag).strip("_ ")

    ext = extension.lower().strip()
    if not ext.startswith("."):
        ext = f".{ext}"

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = f"{safe_stem}_{safe_tag}_{timestamp}{ext}" if safe_tag else f"{safe_stem}_{timestamp}{ext}"
    candidate = out_folder / base_name

    counter = 1
    while candidate.exists():
        alt_name = (
            f"{safe_stem}_{safe_tag}_{timestamp}_{counter}{ext}"
            if safe_tag
            else f"{safe_stem}_{timestamp}_{counter}{ext}"
        )
        candidate = out_folder / alt_name
        counter += 1

    return candidate


def write_formatted_excel(
    sheets_data: pd.DataFrame | dict[str, pd.DataFrame],
    output_path: Path | str,
    highlight_cells: dict[str, set[tuple[int, int]]] | None = None,
    progress_callback: Callable[[int, str], None] | None = None,
    cancel_token: CancelToken | None = None,
) -> Path:
    """
    Escribe uno o varios DataFrames en un archivo .xlsx con formato corporativo:
    - Encabezados en negrita con fondo azul oscuro (#1B365D) y letra blanca
    - Autofiltro activo en todas las hojas
    - Fila superior inmovilizada (freeze panes en A2)
    - Ancho de columnas autoajustado según contenido
    - Formato numérico y de fechas en celdas correspondientes
    - Resaltado opcional de celdas específicas (fila_0_based, col_0_based)
    """
    target = clean_file_path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)

    if isinstance(sheets_data, pd.DataFrame):
        sheets_dict = {"Datos": sheets_data}
    else:
        sheets_dict = sheets_data

    wb = openpyxl.Workbook()
    # Eliminar la hoja por defecto luego de crear las reales
    default_ws = wb.active

    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=False)

    thin_side = Side(border_style="thin", color="D1D5DB")
    header_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    diff_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    diff_font = Font(name="Calibri", size=11, bold=True, color="92400E")

    total_sheets = max(1, len(sheets_dict))
    for idx, (raw_sheet_name, df) in enumerate(sheets_dict.items()):
        if cancel_token:
            cancel_token.raise_if_cancelled()

        # Limpiar nombre de hoja (máximo 31 caracteres en Excel, sin caracteres prohibidos)
        clean_title = "".join(c if c not in '[]:*?/\\' else "_" for c in str(raw_sheet_name))[:31] or f"Hoja{idx + 1}"
        ws = wb.create_sheet(title=clean_title)

        cols = [str(c) for c in df.columns]
        if not cols:
            cols = ["(Sin columnas)"]
            ws.append(cols)
            continue

        # Escribir encabezados
        ws.append(cols)
        ws.row_dimensions[1].height = 24

        for col_idx in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = header_border

        # Escribir filas de datos
        total_rows = len(df)
        sheet_highlights = (highlight_cells or {}).get(raw_sheet_name, set())

        for row_idx, row_tuple in enumerate(df.itertuples(index=False, name=None)):
            if row_idx % 2000 == 0 and cancel_token:
                cancel_token.raise_if_cancelled()

            excel_row_num = row_idx + 2
            cleaned_row: list[Any] = []
            for val in row_tuple:
                if pd.isna(val) or val is None:
                    cleaned_row.append(None)
                elif isinstance(val, (pd.Timestamp, datetime, date)):
                    cleaned_row.append(val)
                elif isinstance(val, (int, float, bool)):
                    cleaned_row.append(val)
                else:
                    cleaned_row.append(str(val))

            ws.append(cleaned_row)

            # Aplicar formatos de celda (fechas, números y resaltados)
            for col_idx, val in enumerate(cleaned_row):
                cell = ws.cell(row=excel_row_num, column=col_idx + 1)
                if isinstance(val, (pd.Timestamp, datetime)):
                    if val.hour == 0 and val.minute == 0 and val.second == 0:
                        cell.number_format = "yyyy-mm-dd"
                    else:
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(val, date) and not isinstance(val, datetime):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(val, float):
                    cell.number_format = "#,##0.00"
                elif isinstance(val, int) and not isinstance(val, bool):
                    cell.number_format = "#,##0"

                if (row_idx, col_idx) in sheet_highlights:
                    cell.fill = diff_fill
                    cell.font = diff_font

            if progress_callback and total_rows > 0 and row_idx % 2500 == 0:
                base_pct = 85 + int(((idx + (row_idx / total_rows)) / total_sheets) * 14)
                progress_callback(min(99, base_pct), f"Formateando hoja '{clean_title}' ({row_idx:,}/{total_rows:,})...")

        # Inmovilizar fila superior (Freeze Panes)
        ws.freeze_panes = "A2"

        # Activar Autofiltro
        last_col_letter = get_column_letter(len(cols))
        last_row_num = max(1, total_rows + 1)
        ws.auto_filter.ref = f"A1:{last_col_letter}{last_row_num}"

        # Ajustar ancho de columnas basado en encabezado + muestra de las primeras 200 filas
        sample_df = df.head(200)
        for col_idx, col_name in enumerate(cols, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(col_name))
            if not sample_df.empty and col_idx - 1 < len(sample_df.columns):
                series = sample_df.iloc[:, col_idx - 1].dropna().astype(str)
                if not series.empty:
                    max_val_len = int(series.map(len).max())
                    max_len = max(max_len, max_val_len)
            adjusted_width = max(12, min(50, max_len + 4))
            ws.column_dimensions[col_letter].width = adjusted_width

    if default_ws in wb.worksheets and len(wb.worksheets) > 1:
        wb.remove(default_ws)

    wb.save(target)
    wb.close()
    if progress_callback:
        progress_callback(100, f"Archivo Excel generado: {target.name}")
    return target


def write_csv_file(
    df: pd.DataFrame,
    output_path: Path | str,
    delimiter: str = ",",
    encoding: str = "utf-8-sig",
    progress_callback: Callable[[int, str], None] | None = None,
    cancel_token: CancelToken | None = None,
) -> Path:
    """
    Exporta un DataFrame a formato CSV con delimitador y codificación configurables
    (por defecto utf-8-sig para que Excel abra bien los acentos y la ñ).
    """
    target = clean_file_path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)

    eff_delim = normalize_delimiter_input(delimiter) or ","
    eff_enc = normalize_encoding_input(encoding) or "utf-8-sig"

    if cancel_token:
        cancel_token.raise_if_cancelled()

    if progress_callback:
        progress_callback(90, f"Escribiendo archivo CSV ({eff_enc}, sep='{eff_delim}')...")

    df.to_csv(
        target,
        sep=eff_delim,
        encoding=eff_enc,
        index=False,
    )

    if progress_callback:
        progress_callback(100, f"Archivo CSV generado: {target.name}")

    return target


def export_dataframe(
    df: pd.DataFrame,
    source_path_or_name: Path | str,
    suffix_tag: str,
    output_format: str = "xlsx",
    output_dir: Path | str | None = None,
    csv_delimiter: str = ",",
    csv_encoding: str = "utf-8-sig",
    sheet_name: str = "Datos",
    progress_callback: Callable[[int, str], None] | None = None,
    cancel_token: CancelToken | None = None,
) -> Path:
    """
    Exporta un DataFrame a Excel (.xlsx con formato) o CSV según la selección del usuario,
    generando un nombre de archivo único con timestamp que jamás sobrescribe el original.
    """
    fmt = output_format.lower().strip().lstrip(".")
    if fmt not in {"xlsx", "csv"}:
        fmt = "xlsx"

    out_path = build_output_path(
        source_path_or_name=source_path_or_name,
        suffix_tag=suffix_tag,
        extension=fmt,
        output_dir=output_dir,
    )

    if fmt == "xlsx":
        return write_formatted_excel(
            sheets_data={sheet_name: df},
            output_path=out_path,
            progress_callback=progress_callback,
            cancel_token=cancel_token,
        )
    return write_csv_file(
        df=df,
        output_path=out_path,
        delimiter=csv_delimiter,
        encoding=csv_encoding,
        progress_callback=progress_callback,
        cancel_token=cancel_token,
    )

